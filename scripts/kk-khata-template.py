"""
The free udhar khata Excel template on mdgservices.in/khata-kavach/credit/udhar-khata-format/,
in English and Hindi: public/kk/downloads/udhar-khata-format.xlsx and udhar-khata-format-hindi.xlsx.

    python3 scripts/kk-khata-template.py

Three sheets:
  - How to use: short steps (Protected View first) and the colour legend.
  - Customers: one row per customer. You type the name, phone, credit limit, credit days and
    the balance from the old book; the sheet works out total udhar, total received, the balance
    now and the credit left.
  - Khata: one row for every credit sale and every payment, the day it happens. You pick the
    customer from the list; the sheet works out that customer's balance after the line.

Yellow cells are typed, grey cells are formulas. The first rows of Customers and Khata are a
made-up example, italic only while their Note still says so. A check cell on Customers counts
Khata lines whose customer is not on the list. Red and orange follow the limit only. Plain Excel-2007 functions only (SUMIFS, INDEX,
MATCH, IFERROR, N), so it opens the same in Excel, Google Sheets, LibreOffice and phone apps.
"""
from pathlib import Path
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from math import ceil
from openpyxl.workbook.properties import CalcProperties

OUT = Path(__file__).resolve().parent.parent / "public" / "kk" / "downloads"

CUST_ROWS = 200   # customers, rows 4..203
KHATA_ROWS = 3000  # lines, rows 4..3003 (customer totals read the whole column, so nothing is ever left out)
FIRST = 4

PINE = "0F4A3D"
INPUT = PatternFill("solid", fgColor="FFF6D5")   # yellow: you type
CALC = PatternFill("solid", fgColor="ECEFEC")    # grey: works out by itself
HEAD = PatternFill("solid", fgColor=PINE)
OVER = PatternFill("solid", fgColor="FFC7CE")    # over the limit (with dark red bold text)
NEAR = PatternFill("solid", fgColor="F8CBAD")    # under a fifth of the limit left: orange (with dark bold text)
thin = Side(style="thin", color="C9D2CE")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
F = "Arial"
OVER_FONT = Font(name="Arial", bold=True, color="9C0006")
NEAR_FONT = Font(name="Arial", bold=True, color="7F3F00")
EXAMPLE_FONT = Font(name="Arial", italic=True, color="5B6964")
MONEY = "#,##0;-#,##0"  # Excel draws the grouping the computer's region uses (12,34,567 in India)

T = {
    "en": {
        "file": "udhar-khata-format.xlsx",
        "how": "How to use", "cust": "Customers", "khata": "Khata",
        "title_how": "Udhar khata: a template from MDG Services",
        "steps": [
            "If Excel shows a yellow 'Protected View' bar at the top, press Enable Editing. The grey columns fill in after that.",
            "1. Customers sheet: one row for each customer. Type the name, phone, credit limit (Rs), credit days (for your own reference) and the balance from your old book. Keep every name different, and do not rename a customer after you have written lines for them.",
            "2. Khata sheet: one row for every credit sale and every payment, written the day it happens. Pick the customer from the list in column B.",
            "3. A sale: fill Udhar given (Rs). A payment: fill Money received (Rs), how it was paid and the UTR or cheque number.",
            "4. Never type over an old row. If a line is wrong, add a new line with the same amount as a minus figure in the same column (Udhar given for a wrong sale, Money received for a wrong payment), write why in Note, then add the right line.",
            "5. The grey columns work out by themselves: each customer's balance, credit left, and the balance after every line. Do not type in them.",
            "6. On the Customers sheet a row turns red when the balance is over the limit, and orange when less than a fifth of the limit is left. On the Khata sheet the balance turns red on a line that took the customer over the limit. The colours follow the limit only: the sheet does not mark late payments.",
            "7. The box at the top of the Customers sheet counts Khata lines whose customer is not on the list. If it shows more than 0, correct those names.",
            "8. The sheet holds 200 customers and 3,000 lines. When the lines run out, start a fresh copy and type each customer's balance now as their balance from the old book.",
            "9. Each month, filter the Khata sheet by the customer's name and send them the lines and the balance.",
            "10. The rows in italics are a made-up example. Delete them, or type over every cell including the Note, before you start.",
        ],
        "legend": "Yellow cells: you type. Grey cells: worked out by the sheet, do not type in them.",
        "swatch_in": "Yellow cells: you type",
        "swatch_calc": "Grey cells: worked out by the sheet. Do not type in them.",
        "note": "General information. Check your own figures. The names in the example rows are made up.",
        "by": "MDG Services, a brand of VRUOOM Private Limited · mdgservices.in/khata-kavach/credit/udhar-khata-format/",
        "kk": "Khata Kavach, our app, keeps credit sales and payments in one shared book: by default the customer sends a request before goods go out, and both sides see each line as soon as the phone is online. mdgservices.in/khata-kavach",
        "title_cust": "Customers: limits and balances",
        "title_khata": "Khata: every credit sale and every payment",
        "check": "Khata lines whose customer is not on this list:",
        "cust_cols": ["Customer name", "Phone", "Credit limit (Rs)", "Credit days (for reference)", "Balance from old book (Rs)", "Total udhar given (Rs)", "Total received (Rs)", "Balance now (Rs)", "Credit left (Rs)", "Note"],
        "khata_cols": ["Date", "Customer (pick from list)", "Vehicle no.", "Item", "Quantity", "Unit", "Udhar given (Rs)", "Money received (Rs)", "How paid", "UTR / cheque no.", "Slip / bill no.", "Who served", "Customer's balance after this line (Rs)", "Note"],
        "units": ["Litre", "Kg", "Tonne", "Bag", "Piece", "Box", "Cylinder", "Trip", "Metre", "Amount only"],
        "paid": ["Cash", "UPI", "Bank", "Cheque"],
        "pick_err": "Pick the customer from the list. Add a new customer on the Customers sheet first.",
        "num_err": "Type a number only. Use a minus sign to correct a line.",
        "ex_key": "Example",
        "ex_cust": ["Sharma Roadlines (made-up)", "98XXXXXX00", 200000, 15, 20000, None, None, None, None, "Example row: delete it or type over it"],
        "ex_lines": [
            [date(2026, 10, 2), "UP00 SH 1005", "Diesel", 150, "Litre", 13500, None, None, None, "S-101", "Staff A", "Example row"],
            [date(2026, 10, 4), "UP00 SH 1006", "Diesel", 200, "Litre", 18000, None, None, None, "S-102", "Staff A", "Example row"],
            [date(2026, 10, 10), None, None, None, None, None, 30000, "UPI", "UTR 0000XXXX1234", None, "Owner", "Example row"],
        ],
    },
    "hi": {
        "file": "udhar-khata-format-hindi.xlsx",
        "how": "कैसे भरें", "cust": "ग्राहक", "khata": "खाता",
        "title_how": "उधार खाता: MDG Services का फ़ॉर्मेट",
        "steps": [
            "अगर Excel ऊपर पीली पट्टी में 'Protected View' दिखाए, तो 'Enable Editing' दबाएँ। उसके बाद स्लेटी कॉलम भरते हैं।",
            "1. ग्राहक शीट: हर ग्राहक की एक लाइन। नाम, फ़ोन, उधार की सीमा (₹), उधार के दिन (याद के लिए) और पुरानी बही का बकाया लिखें। हर ग्राहक का नाम अलग रखें, और लाइनें लिख देने के बाद नाम न बदलें।",
            "2. खाता शीट: हर उधार की बिक्री और हर पेमेंट की एक लाइन, उसी दिन लिखें। कॉलम B में लिस्ट से ग्राहक चुनें।",
            "3. बिक्री: 'उधार दिया (₹)' भरें। पेमेंट: 'पैसे मिले (₹)', कैसे मिले, और UTR या चेक नंबर भरें।",
            "4. पुरानी लाइन पर कभी दोबारा न लिखें। लाइन गलत हो तो नई लाइन जोड़ें, उसी खाने में उतनी ही रकम माइनस में (गलत बिक्री हो तो 'उधार दिया' में, गलत पेमेंट हो तो 'पैसे मिले' में), 'नोट' में वजह लिखें, फिर सही लाइन जोड़ें।",
            "5. स्लेटी कॉलम अपने आप भरते हैं: हर ग्राहक का बकाया, बची हुई सीमा, और हर लाइन के बाद का बकाया। इनमें कुछ न लिखें।",
            "6. ग्राहक शीट पर बकाया सीमा से ऊपर जाए तो लाइन लाल हो जाती है, और सीमा का पाँचवाँ हिस्सा भी न बचे तो नारंगी। खाता शीट पर जिस लाइन से ग्राहक सीमा से ऊपर गया, उसका बकाया लाल दिखता है। रंग सिर्फ़ सीमा देखते हैं: देर से पेमेंट पर शीट कोई निशान नहीं लगाती।",
            "7. ग्राहक शीट के ऊपर एक खाना गिनता है कि खाता शीट की कितनी लाइनों का ग्राहक लिस्ट में नहीं है। 0 से ज़्यादा दिखे तो उन लाइनों के नाम ठीक करें।",
            "8. शीट में 200 ग्राहक और 3,000 लाइनें आती हैं। लाइनें भर जाएँ तो नई कॉपी शुरू करें और हर ग्राहक का अभी का बकाया उसकी 'पुरानी बही का बकाया' में लिखें।",
            "9. हर महीने खाता शीट को ग्राहक के नाम से फ़िल्टर करें और उसे लाइनें और बकाया भेजें।",
            "10. तिरछे अक्षरों वाली लाइनें काल्पनिक उदाहरण हैं। शुरू करने से पहले उन्हें हटा दें, या 'नोट' समेत हर खाने में अपना लिखें।",
        ],
        "legend": "पीले खाने: आप भरें। स्लेटी खाने: शीट ख़ुद भरती है, इनमें कुछ न लिखें।",
        "swatch_in": "पीले खाने: आप भरें",
        "swatch_calc": "स्लेटी खाने: शीट ख़ुद भरती है। इनमें कुछ न लिखें।",
        "note": "यह सामान्य जानकारी है। अपने आँकड़े ख़ुद जाँचें। उदाहरण वाली लाइनों के नाम काल्पनिक हैं।",
        "by": "MDG Services, VRUOOM Private Limited का ब्रांड · mdgservices.in/hi/khata-kavach/credit/udhar-khata-format/",
        "kk": "हमारा ऐप खाता कवच उधार की बिक्री और पेमेंट एक साझा खाते में रखता है: शुरुआती सेटिंग में माल जाने से पहले ग्राहक रिक्वेस्ट भेजता है, और फ़ोन के नेट पर आते ही हर लाइन दोनों तरफ़ दिखती है। mdgservices.in/hi/khata-kavach",
        "title_cust": "ग्राहक: उधार की सीमा और बकाया",
        "title_khata": "खाता: हर उधार की बिक्री और हर पेमेंट",
        "check": "खाता शीट की लाइनें जिनका ग्राहक इस लिस्ट में नहीं:",
        "cust_cols": ["ग्राहक का नाम", "फ़ोन", "उधार की सीमा (₹)", "उधार के दिन (याद के लिए)", "पुरानी बही का बकाया (₹)", "कुल उधार दिया (₹)", "कुल पैसे मिले (₹)", "अभी का बकाया (₹)", "बची हुई सीमा (₹)", "नोट"],
        "khata_cols": ["तारीख़", "ग्राहक (लिस्ट से चुनें)", "गाड़ी नंबर", "माल", "मात्रा", "यूनिट", "उधार दिया (₹)", "पैसे मिले (₹)", "कैसे मिले", "UTR / चेक नंबर", "पर्ची / बिल नंबर", "स्टाफ़ का नाम", "इस लाइन के बाद ग्राहक का बकाया (₹)", "नोट"],
        "units": ["लीटर", "किलो", "टन", "बोरी", "पीस", "बॉक्स", "सिलेंडर", "ट्रिप", "मीटर", "कोई यूनिट नहीं (सिर्फ़ रकम)"],
        "paid": ["नकद", "UPI", "बैंक", "चेक"],
        "pick_err": "ग्राहक लिस्ट से चुनें। नया ग्राहक पहले ग्राहक शीट में जोड़ें।",
        "num_err": "सिर्फ़ अंक भरें। लाइन सुधारने के लिए माइनस लगाएँ।",
        "ex_key": "उदाहरण",
        "ex_cust": ["Sharma Roadlines (काल्पनिक)", "98XXXXXX00", 200000, 15, 20000, None, None, None, None, "उदाहरण की लाइन: इसे हटाएँ या इस पर अपना लिखें"],
        "ex_lines": [
            [date(2026, 10, 2), "UP00 SH 1005", "डीज़ल", 150, "लीटर", 13500, None, None, None, "S-101", "स्टाफ़ A", "उदाहरण की लाइन"],
            [date(2026, 10, 4), "UP00 SH 1006", "डीज़ल", 200, "लीटर", 18000, None, None, None, "S-102", "स्टाफ़ A", "उदाहरण की लाइन"],
            [date(2026, 10, 10), None, None, None, None, None, 30000, "UPI", "UTR 0000XXXX1234", None, "मालिक", "उदाहरण की लाइन"],
        ],
    },
}


def ref(sheet: str) -> str:
    """A sheet name as a formula reference: always quoted (Hindi names and spaces need it)."""
    return "'" + sheet.replace("'", "''") + "'"


def title(ws, text: str, sub: str) -> None:
    ws["A1"] = text
    ws["A1"].font = Font(name=F, size=14, bold=True, color=PINE)
    ws["A2"] = sub
    ws["A2"].font = Font(name=F, size=10, color="43524D")
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[3].height = 34


def header(ws, cols: list[str], widths: list[int]) -> None:
    for i, (name, w) in enumerate(zip(cols, widths), start=1):
        c = ws.cell(row=3, column=i, value=name)
        c.font = Font(name=F, size=10, bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BOX
        ws.column_dimensions[c.column_letter].width = w


def build(lang: str) -> Path:
    t = T[lang]
    wb = Workbook()
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    per_line = 78 if lang == "en" else 62  # characters a How-to-use line holds at width 70

    # ── How to use ──
    how = wb.active
    how.title = t["how"]
    how["A1"] = t["title_how"]
    how["A1"].font = Font(name=F, size=15, bold=True, color=PINE)
    r = 3
    for i, s in enumerate(t["steps"]):
        c = how.cell(row=r, column=1, value=s)
        c.font = Font(name=F, size=11, bold=(i == 0))
        c.alignment = Alignment(wrap_text=True, vertical="top")
        how.row_dimensions[r].height = 17 * ceil(len(s) / per_line) + 6
        r += 1
    r += 1
    # the legend, drawn in the two colours it names
    for text, fill in ((t["swatch_in"], INPUT), (t["swatch_calc"], CALC)):
        c = how.cell(row=r, column=1, value=text)
        c.font = Font(name=F, size=11, bold=True)
        c.fill = fill
        c.border = BOX
        r += 1
    r += 1
    for key in ("note", "kk", "by"):
        c = how.cell(row=r, column=1, value=t[key])
        c.font = Font(name=F, size=9 if key == "by" else 10, color="43524D")
        c.alignment = Alignment(wrap_text=True, vertical="top")
        how.row_dimensions[r].height = 15 * ceil(len(t[key]) / per_line) + 6
        r += 1
    how.column_dimensions["A"].width = 70

    cs, ks = t["cust"], t["khata"]
    last_c, last_k = FIRST + CUST_ROWS - 1, FIRST + KHATA_ROWS - 1
    k, c_ = ref(ks), ref(cs)
    ex = t["ex_key"]

    # ── Customers ──
    cu = wb.create_sheet(cs)
    title(cu, t["title_cust"], t["legend"])
    header(cu, t["cust_cols"], [30, 15, 15, 13, 18, 18, 17, 16, 16, 30])
    # the name check: Khata lines whose customer is not on this list
    cu["F2"] = t["check"]
    cu["F2"].font = Font(name=F, size=10, bold=True, color="43524D")
    cu["J2"] = f'=SUMPRODUCT(({k}!$B${FIRST}:$B${last_k}<>"")*(COUNTIF($A${FIRST}:$A${last_c},{k}!$B${FIRST}:$B${last_k})=0))'
    cu["J2"].font = Font(name=F, size=11, bold=True)
    cu["J2"].fill = CALC
    cu["J2"].border = BOX
    cu["J2"].alignment = Alignment(horizontal="left")
    cu.conditional_formatting.add("J2", FormulaRule(formula=["N($J$2)>0"], fill=OVER, font=OVER_FONT))
    for row in range(FIRST, last_c + 1):
        for col in range(1, 11):
            c = cu.cell(row=row, column=col)
            c.border = BOX
            c.font = Font(name=F, size=10)
            c.fill = CALC if 6 <= col <= 9 else INPUT
        # whole columns: a line typed below the formatted rows still counts
        cu.cell(row=row, column=6, value=f'=IF(A{row}="","",SUMIFS({k}!$G:$G,{k}!$B:$B,A{row}))')
        cu.cell(row=row, column=7, value=f'=IF(A{row}="","",SUMIFS({k}!$H:$H,{k}!$B:$B,A{row}))')
        cu.cell(row=row, column=8, value=f'=IF(A{row}="","",N(E{row})+F{row}-G{row})')
        cu.cell(row=row, column=9, value=f'=IF(OR(A{row}="",C{row}=""),"",C{row}-H{row})')
        for col in (3, 5, 6, 7, 8, 9):
            cu.cell(row=row, column=col).number_format = MONEY
    for col, v in enumerate(t["ex_cust"], start=1):
        if v is not None:
            cu.cell(row=FIRST, column=col, value=v)
    rng = f"A{FIRST}:J{last_c}"
    cu.conditional_formatting.add(rng, FormulaRule(formula=[f'AND($A{FIRST}<>"",N($C{FIRST})>0,N($H{FIRST})>N($C{FIRST}))'], fill=OVER, font=OVER_FONT, stopIfTrue=True))
    cu.conditional_formatting.add(rng, FormulaRule(formula=[f'AND($A{FIRST}<>"",N($C{FIRST})>0,N($I{FIRST})<N($C{FIRST})/5)'], fill=NEAR, font=NEAR_FONT, stopIfTrue=True))
    # the example row is italic only while its Note still says it is an example
    cu.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT($J{FIRST},{len(ex)})="{ex}"'], font=EXAMPLE_FONT))
    cu.freeze_panes = f"B{FIRST}"
    cu.auto_filter.ref = f"A3:J{last_c}"

    # ── Khata ──
    kh = wb.create_sheet(ks)
    title(kh, t["title_khata"], t["legend"])
    header(kh, t["khata_cols"], [12, 28, 15, 14, 10, 13, 15, 16, 11, 20, 13, 14, 20, 30])
    for row in range(FIRST, last_k + 1):
        for col in range(1, 15):
            cell = kh.cell(row=row, column=col)
            cell.border = BOX
            cell.font = Font(name=F, size=10)
            cell.fill = CALC if col == 13 else INPUT
        kh.cell(row=row, column=1).number_format = "dd-mm-yyyy"
        for col in (7, 8, 13):
            kh.cell(row=row, column=col).number_format = MONEY
        kh.cell(row=row, column=13, value=(
            f'=IF(OR(B{row}="",AND(G{row}="",H{row}="")),"",'
            f'IFERROR(N(INDEX({c_}!$E${FIRST}:$E${last_c},MATCH(B{row},{c_}!$A${FIRST}:$A${last_c},0))),0)'
            f'+SUMIFS($G${FIRST}:G{row},$B${FIRST}:B{row},B{row})-SUMIFS($H${FIRST}:H{row},$B${FIRST}:B{row},B{row}))'
        ))
    for i, line in enumerate(t["ex_lines"]):
        row = FIRST + i
        vals = [line[0], t["ex_cust"][0]] + line[1:-1] + [None, line[-1]]  # M, the balance, stays a formula
        for col, v in enumerate(vals, start=1):
            if v is not None:
                kh.cell(row=row, column=col, value=v)
    # pick lists, and numbers only in the two money columns
    dv_c = DataValidation(type="list", formula1=f"{c_}!$A${FIRST}:$A${last_c}", allow_blank=True, showErrorMessage=True, error=t["pick_err"])
    dv_u = DataValidation(type="list", formula1='"' + ",".join(t["units"]) + '"', allow_blank=True)
    dv_p = DataValidation(type="list", formula1='"' + ",".join(t["paid"]) + '"', allow_blank=True)
    dv_n = DataValidation(type="decimal", operator="between", formula1="-1000000000000", formula2="1000000000000", allow_blank=True, showErrorMessage=True, error=t["num_err"])
    for dv, cols in ((dv_c, "B"), (dv_u, "F"), (dv_p, "I"), (dv_n, "GH")):
        kh.add_data_validation(dv)
        for col in cols:
            dv.add(f"{col}{FIRST}:{col}{last_k}")
    limit = f'N(INDEX({c_}!$C${FIRST}:$C${last_c},MATCH($B{FIRST},{c_}!$A${FIRST}:$A${last_c},0)))'
    kh.conditional_formatting.add(
        f"M{FIRST}:M{last_k}",
        FormulaRule(formula=[f'IFERROR(AND($M{FIRST}<>"",{limit}>0,N($M{FIRST})>{limit}),FALSE)'], fill=OVER, font=OVER_FONT),
    )
    kh.conditional_formatting.add(f"A{FIRST}:N{last_k}", FormulaRule(formula=[f'LEFT($N{FIRST},{len(ex)})="{ex}"'], font=EXAMPLE_FONT))
    kh.freeze_panes = f"B{FIRST}"
    kh.auto_filter.ref = f"A3:N{last_k}"

    for ws in (cu, kh):
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = "3:3"

    wb.active = 0
    OUT.mkdir(parents=True, exist_ok=True)
    out = OUT / t["file"]
    wb.save(out)
    return out


if __name__ == "__main__":
    for lang in ("en", "hi"):
        print("wrote", build(lang))
