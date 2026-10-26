"""
The customer statement-of-account Excel sheet on
mdgservices.in/khata-kavach/getting-paid/customer-statement-format/, in English and Hindi:
public/kk/downloads/customer-statement.xlsx and customer-statement-hindi.xlsx.

    python3 scripts/kk-statement.py

Two sheets:
  - How to use: short steps (Protected View first) and the colour legend.
  - Statement: the business, the customer, the period and the opening balance at the top (yellow,
    typed); one row for every sale and every payment in the period (yellow, typed); the sheet
    works out the balance after every line, the totals and the closing balance (grey, formulas).
    A confirmation line and two signature lines sit under the table, so the sheet prints as the
    statement to send.

Yellow cells are typed, grey cells are formulas. The first rows are a made-up example, italic
only while their Details still say so. Plain Excel-2007 functions only (IF, AND, N, SUM), so it
opens the same in Excel, Google Sheets, LibreOffice and phone apps. The printable A4 page is
made by scripts/kk-statement.mjs.
"""
from pathlib import Path
from datetime import date
from math import ceil
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.properties import CalcProperties

OUT = Path(__file__).resolve().parent.parent / "public" / "kk" / "downloads"

LINES = 50          # sale and payment rows
HEAD_ROW = 9        # the column headings
OPEN_ROW = 10       # the "Opening balance" line
FIRST = 11          # first sale or payment row
LAST = FIRST + LINES - 1   # 60
TOT = LAST + 1      # 61: totals and closing balance

PINE = "0F4A3D"
INPUT = PatternFill("solid", fgColor="FFF6D5")   # yellow: you type
CALC = PatternFill("solid", fgColor="ECEFEC")    # grey: works out by itself
HEAD = PatternFill("solid", fgColor=PINE)
thin = Side(style="thin", color="C9D2CE")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
F = "Arial"
EXAMPLE_FONT = Font(name="Arial", italic=True, color="5B6964")
MONEY = "#,##0;-#,##0"  # Excel draws the grouping the computer's region uses (12,34,567 in India)

T = {
    "en": {
        "file": "customer-statement.xlsx",
        "how": "How to use", "sheet": "Statement",
        "title_how": "Statement of account: a template from MDG Services",
        "steps": [
            "If Excel shows a yellow 'Protected View' bar at the top, press Enable Editing. The grey cells fill in after that.",
            "1. At the top of the Statement sheet, type your business name, the customer's name, the statement date, the period it covers (from and to) and the opening balance.",
            "2. Opening balance: what the customer owed you at the start of the period. For a second statement, it is the closing balance of the one before.",
            "3. One row for every sale and every payment in the period, in date order. Type the date, the bill or slip number, the vehicle or who collected, and the details (item and quantity for a sale; how it was paid and the UTR or cheque number for a payment).",
            "4. A sale goes in Udhar given (Rs). A payment goes in Money received (Rs). Never both on one row.",
            "5. The grey cells work out by themselves: the balance after every line (opening balance + all udhar given so far - all money received so far), the totals, and the closing balance. Do not type in them.",
            "6. The closing balance, at the top and at the foot of the table, is what the customer owes you at the end of the period.",
            "7. Mistakes: before you send the statement, fix them on the row. After you have sent it, do not change the old row. Add a new row with the same amount as a minus figure in the same column, write why in Details, then add the right row. Both copies then tell the same story.",
            "8. The sheet holds 50 rows. If the period has more, make a second statement and start it with the closing balance of the first as its opening balance.",
            "9. Print it on A4 or save it as a PDF, and send it from your own WhatsApp. Ask the customer to check it and to confirm it or tell you what looks wrong within 7 days.",
            "10. The rows in italics are a made-up example. Delete them, or type over every cell including the Details, before you start.",
        ],
        "legend": "Yellow cells: you type. Grey cells: worked out by the sheet, do not type in them.",
        "swatch_in": "Yellow cells: you type",
        "swatch_calc": "Grey cells: worked out by the sheet. Do not type in them.",
        "note": "General information. Check your own figures. The names in the example rows are made up.",
        "by": "MDG Services, a brand of VRUOOM Private Limited · mdgservices.in/khata-kavach/getting-paid/customer-statement-format/",
        "kk": "Khata Kavach, our app, keeps credit sales and payments in one shared book: by default the customer sends a request before goods go out, and both sides see each line as soon as the phone is online. mdgservices.in/khata-kavach",
        "title": "Statement of account",
        "sub": "Fill the yellow cells. The grey cells work themselves out.",
        "business": "Business", "customer": "Customer", "opening": "Opening balance (Rs)", "closing": "Closing balance (Rs)",
        "stmt_date": "Statement date", "from": "Period from", "to": "Period to",
        "cols": ["Date", "Bill / slip no.", "Vehicle no. / collected by", "Details (item and quantity, or how paid and UTR / cheque no.)", "Udhar given (Rs)", "Money received (Rs)", "Balance (Rs)"],
        "open_line": "Opening balance", "totals": "Totals", "close_line": "Closing balance",
        "confirm": "Please check this statement. If it is correct, sign below. If anything looks wrong, tell us within 7 days and we will look at it together.",
        "sign_biz": "For the business: sign and date", "sign_cust": "Customer's confirmation: name, sign and date",
        "num_err": "Type a number only. Use a minus sign to correct a line.",
        "ex_key": "Example",
        "ex_biz": "Sharma Traders (made-up)", "ex_cust": "Gupta Transport (made-up)",
        "ex_lines": [
            [date(2026, 10, 2), "S-101", "UP00 SH 1005", "Example: diesel, 150 litre", 13500, None],
            [date(2026, 10, 4), "S-102", "UP00 SH 1006", "Example: diesel, 200 litre", 18000, None],
            [date(2026, 10, 10), None, None, "Example: UPI, UTR 0000XXXX1234", None, 30000],
        ],
    },
    "hi": {
        "file": "customer-statement-hindi.xlsx",
        "how": "कैसे भरें", "sheet": "स्टेटमेंट",
        "title_how": "खाते का स्टेटमेंट: MDG Services का फ़ॉर्मेट",
        "steps": [
            "अगर Excel ऊपर पीली पट्टी में 'Protected View' दिखाए, तो 'Enable Editing' दबाएँ। उसके बाद स्लेटी खाने भरते हैं।",
            "1. स्टेटमेंट शीट के ऊपर अपने कारोबार का नाम, ग्राहक का नाम, स्टेटमेंट की तारीख़, कौन-सी अवधि का है (कब से, कब तक) और शुरू का बकाया लिखें।",
            "2. शुरू का बकाया: अवधि की शुरुआत में ग्राहक पर आपका जितना बाकी था। दूसरा स्टेटमेंट बना रहे हों तो पिछले स्टेटमेंट का आख़िरी बकाया।",
            "3. अवधि की हर बिक्री और हर पेमेंट की एक लाइन, तारीख़ के क्रम में। तारीख़, बिल या पर्ची नंबर, गाड़ी या माल लेने वाला, और ब्योरा लिखें (बिक्री में माल और मात्रा; पेमेंट में कैसे मिला और UTR या चेक नंबर)।",
            "4. बिक्री 'उधार दिया (₹)' में जाती है। पेमेंट 'पैसे मिले (₹)' में जाता है। एक लाइन में दोनों नहीं।",
            "5. स्लेटी खाने अपने आप भरते हैं: हर लाइन के बाद का बकाया (शुरू का बकाया + अब तक का सारा उधार दिया - अब तक के सारे पैसे मिले), जोड़, और आख़िरी बकाया। इनमें कुछ न लिखें।",
            "6. आख़िरी बकाया, ऊपर और टेबल के नीचे, वह रकम है जो अवधि के अंत में ग्राहक पर आपकी बाकी है।",
            "7. गलती: स्टेटमेंट भेजने से पहले उसे उसी लाइन पर ठीक कर लें। भेजने के बाद पुरानी लाइन न बदलें। उसी खाने में उतनी ही रकम माइनस में डालकर नई लाइन जोड़ें, 'ब्योरा' में वजह लिखें, फिर सही लाइन जोड़ें। तब दोनों कॉपियों में एक ही कहानी रहती है।",
            "8. शीट में 50 लाइनें आती हैं। अवधि में ज़्यादा हों तो दूसरा स्टेटमेंट बनाएँ और उसमें पहले का आख़िरी बकाया शुरू का बकाया रखें।",
            "9. A4 पर छापें या PDF बनाएँ, और अपने WhatsApp से भेजें। ग्राहक से कहें कि जाँच ले, और 7 दिन के भीतर बताए कि ठीक है या क्या गलत लग रहा है।",
            "10. तिरछे अक्षरों वाली लाइनें काल्पनिक उदाहरण हैं। शुरू करने से पहले उन्हें हटा दें, या 'ब्योरा' समेत हर खाने में अपना लिखें।",
        ],
        "legend": "पीले खाने: आप भरें। स्लेटी खाने: शीट ख़ुद भरती है, इनमें कुछ न लिखें।",
        "swatch_in": "पीले खाने: आप भरें",
        "swatch_calc": "स्लेटी खाने: शीट ख़ुद भरती है। इनमें कुछ न लिखें।",
        "note": "यह सामान्य जानकारी है। अपने आँकड़े ख़ुद जाँचें। उदाहरण वाली लाइनों के नाम काल्पनिक हैं।",
        "by": "MDG Services, VRUOOM Private Limited का ब्रांड · mdgservices.in/hi/khata-kavach/getting-paid/customer-statement-format/",
        "kk": "हमारा ऐप खाता कवच उधार की बिक्री और पेमेंट एक साझा खाते में रखता है: शुरुआती सेटिंग में माल जाने से पहले ग्राहक रिक्वेस्ट भेजता है, और फ़ोन के नेट पर आते ही हर लाइन दोनों तरफ़ दिखती है। mdgservices.in/hi/khata-kavach",
        "title": "खाते का स्टेटमेंट",
        "sub": "पीले खाने भरें। स्लेटी खाने अपने आप भरते हैं।",
        "business": "आपका कारोबार", "customer": "ग्राहक", "opening": "शुरू का बकाया (₹)", "closing": "आख़िरी बकाया (₹)",
        "stmt_date": "स्टेटमेंट की तारीख़", "from": "अवधि: कब से", "to": "अवधि: कब तक",
        "cols": ["तारीख़", "बिल / पर्ची नंबर", "गाड़ी नंबर / माल लेने वाला", "ब्योरा (माल और मात्रा, या कैसे मिले और UTR / चेक नंबर)", "उधार दिया (₹)", "पैसे मिले (₹)", "बकाया (₹)"],
        "open_line": "शुरू का बकाया", "totals": "जोड़", "close_line": "आख़िरी बकाया",
        "confirm": "कृपया यह स्टेटमेंट जाँच लें। ठीक हो तो नीचे दस्तख़त करें। कुछ गलत लगे तो 7 दिन के भीतर हमें बताएँ, हम साथ बैठकर देख लेंगे।",
        "sign_biz": "कारोबार की ओर से: दस्तख़त और तारीख़", "sign_cust": "ग्राहक की पुष्टि: नाम, दस्तख़त और तारीख़",
        "num_err": "सिर्फ़ अंक भरें। लाइन सुधारने के लिए माइनस लगाएँ।",
        "ex_key": "उदाहरण",
        "ex_biz": "Sharma Traders (काल्पनिक)", "ex_cust": "Gupta Transport (काल्पनिक)",
        "ex_lines": [
            [date(2026, 10, 2), "S-101", "UP00 SH 1005", "उदाहरण: डीज़ल, 150 लीटर", 13500, None],
            [date(2026, 10, 4), "S-102", "UP00 SH 1006", "उदाहरण: डीज़ल, 200 लीटर", 18000, None],
            [date(2026, 10, 10), None, None, "उदाहरण: UPI, UTR 0000XXXX1234", None, 30000],
        ],
    },
}


def label(ws, cell: str, text: str) -> None:
    ws[cell] = text
    ws[cell].font = Font(name=F, size=10, bold=True, color="43524D")
    ws[cell].alignment = Alignment(vertical="center")


def field(ws, rng: str, fill, number_format=None) -> None:
    """A merged box: the first cell holds the value, every cell carries the colour and the border."""
    first = rng.split(":")[0]
    ws.merge_cells(rng)
    for row in ws[rng]:
        for c in row:
            c.fill = fill
            c.border = BOX
            c.font = Font(name=F, size=11, bold=(fill is CALC))
    ws[first].alignment = Alignment(horizontal="left", vertical="center")
    if number_format:
        ws[first].number_format = number_format


def build(lang: str) -> Path:
    t = T[lang]
    wb = Workbook()
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    per_line = 78 if lang == "en" else 62  # characters a How-to-use line holds at width 70

    # ── How to use ──
    how = wb.active
    how.title = t["how"]
    how["A1"] = t["title_how"]
    how["A1"].font = Font(name=F, size=15, bold=True, color=PINE)
    r = 3
    for i, s in enumerate(t["steps"]):
        c = how.cell(row=r, column=1, value=s)
        c.font = Font(name=F, size=11, bold=(i == 0))
        c.alignment = Alignment(wrap_text=True, vertical="top")
        how.row_dimensions[r].height = 17 * ceil(len(s) / per_line) + 6
        r += 1
    r += 1
    for text, fill in ((t["swatch_in"], INPUT), (t["swatch_calc"], CALC)):
        c = how.cell(row=r, column=1, value=text)
        c.font = Font(name=F, size=11, bold=True)
        c.fill = fill
        c.border = BOX
        r += 1
    r += 1
    for key in ("note", "kk", "by"):
        c = how.cell(row=r, column=1, value=t[key])
        c.font = Font(name=F, size=9 if key == "by" else 10, color="43524D")
        c.alignment = Alignment(wrap_text=True, vertical="top")
        how.row_dimensions[r].height = 15 * ceil(len(t[key]) / per_line) + 6
        r += 1
    how.column_dimensions["A"].width = 70

    # ── Statement ──
    ws = wb.create_sheet(t["sheet"])
    for col, w in zip("ABCDEFG", [13, 15, 21, 42, 22, 16, 17]):
        ws.column_dimensions[col].width = w
    ws["A1"] = t["title"]
    ws["A1"].font = Font(name=F, size=16, bold=True, color=PINE)
    ws["A2"] = t["sub"]
    ws["A2"].font = Font(name=F, size=10, color="43524D")
    ws.row_dimensions[1].height = 26

    # the top block: left = business, customer, opening and closing; right = dates
    label(ws, "A4", t["business"]);   field(ws, "C4:D4", INPUT)
    label(ws, "A5", t["customer"]);   field(ws, "C5:D5", INPUT)
    label(ws, "A6", t["opening"]);    field(ws, "C6:D6", INPUT, MONEY)
    label(ws, "A7", t["closing"]);    field(ws, "C7:D7", CALC, MONEY)
    label(ws, "E4", t["stmt_date"]);  field(ws, "F4:G4", INPUT, "dd-mm-yyyy")
    label(ws, "E5", t["from"]);       field(ws, "F5:G5", INPUT, "dd-mm-yyyy")
    label(ws, "E6", t["to"]);         field(ws, "F6:G6", INPUT, "dd-mm-yyyy")
    ws["C7"] = f"=G{TOT}"
    for row in (4, 5, 6, 7):
        ws.row_dimensions[row].height = 20

    # column headings
    ws.row_dimensions[HEAD_ROW].height = 44
    for i, name in enumerate(t["cols"], start=1):
        c = ws.cell(row=HEAD_ROW, column=i, value=name)
        c.font = Font(name=F, size=10, bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BOX

    # the opening balance line
    for col in range(1, 8):
        c = ws.cell(row=OPEN_ROW, column=col)
        c.fill = CALC
        c.border = BOX
        c.font = Font(name=F, size=10, bold=True)
    ws.cell(row=OPEN_ROW, column=1, value='=IF(F5="","",F5)').number_format = "dd-mm-yyyy"
    ws.cell(row=OPEN_ROW, column=4, value=t["open_line"])
    ws.cell(row=OPEN_ROW, column=7, value="=N($C$6)").number_format = MONEY

    # the sale and payment rows
    for row in range(FIRST, LAST + 1):
        for col in range(1, 8):
            c = ws.cell(row=row, column=col)
            c.border = BOX
            c.font = Font(name=F, size=10)
            c.fill = CALC if col == 7 else INPUT
        ws.cell(row=row, column=1).number_format = "dd-mm-yyyy"
        ws.cell(row=row, column=4).alignment = Alignment(wrap_text=True, vertical="top")
        for col in (5, 6, 7):
            ws.cell(row=row, column=col).number_format = MONEY
        ws.cell(row=row, column=7, value=(
            f'=IF(AND(E{row}="",F{row}=""),"",N($G${OPEN_ROW})+SUM($E${FIRST}:E{row})-SUM($F${FIRST}:F{row}))'
        ))
    for i, line in enumerate(t["ex_lines"]):
        row = FIRST + i
        for col, v in enumerate(line, start=1):
            if v is not None:
                ws.cell(row=row, column=col, value=v)
    ws["C4"] = t["ex_biz"]
    ws["C5"] = t["ex_cust"]
    ws["C6"] = 20000
    ws["F4"] = date(2026, 11, 1)
    ws["F5"] = date(2026, 10, 1)
    ws["F6"] = date(2026, 10, 31)

    # totals and closing balance
    for col in range(1, 8):
        c = ws.cell(row=TOT, column=col)
        c.fill = CALC
        c.border = BOX
        c.font = Font(name=F, size=10, bold=True)
    ws.cell(row=TOT, column=4, value=f'{t["totals"]} · {t["close_line"]}')
    ws.cell(row=TOT, column=5, value=f"=SUM(E{FIRST}:E{LAST})").number_format = MONEY
    ws.cell(row=TOT, column=6, value=f"=SUM(F{FIRST}:F{LAST})").number_format = MONEY
    ws.cell(row=TOT, column=7, value=f"=N($C$6)+E{TOT}-F{TOT}").number_format = MONEY

    # confirmation and signatures
    r = TOT + 2
    ws.merge_cells(start_row=r, start_column=1, end_row=r + 1, end_column=7)
    ws.cell(row=r, column=1, value=t["confirm"])
    ws.cell(row=r, column=1).font = Font(name=F, size=10)
    ws.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    sig = r + 4
    ws.row_dimensions[sig].height = 26
    for start, end in ((1, 3), (4, 7)):
        for col in range(start, end + 1):
            ws.cell(row=sig, column=col).border = Border(bottom=Side(style="thin", color="43524D"))
    ws.cell(row=sig + 1, column=1, value=t["sign_biz"]).font = Font(name=F, size=9, color="43524D")
    ws.cell(row=sig + 1, column=4, value=t["sign_cust"]).font = Font(name=F, size=9, color="43524D")

    # numbers only in the two money columns (a minus figure corrects a line)
    dv = DataValidation(type="decimal", operator="between", formula1="-1000000000000", formula2="1000000000000", allow_blank=True, showErrorMessage=True, error=t["num_err"])
    ws.add_data_validation(dv)
    dv.add(f"E{FIRST}:F{LAST}")
    dv.add("C6")
    # an example row is italic only while its Details still start with the example word
    ex = t["ex_key"]
    ws.conditional_formatting.add(f"A{FIRST}:G{LAST}", FormulaRule(formula=[f'LEFT($D{FIRST},{len(ex)})="{ex}"'], font=EXAMPLE_FONT))

    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{HEAD_ROW}:{HEAD_ROW}"
    ws.print_area = f"A1:G{sig + 1}"
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5

    wb.active = 0
    OUT.mkdir(parents=True, exist_ok=True)
    out = OUT / t["file"]
    wb.save(out)
    return out


if __name__ == "__main__":
    for lang in ("en", "hi"):
        print("wrote", build(lang))
